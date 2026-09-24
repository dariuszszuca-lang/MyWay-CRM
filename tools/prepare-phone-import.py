#!/usr/bin/env python3
"""Reconcile all seven local attachments. Produces a private, reviewed import payload."""
import argparse, csv, datetime, hashlib, importlib.util, io, json, os, pathlib, re, tempfile, zipfile
from numbers_parser import Document
from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location('preview', pathlib.Path(__file__).with_name('preview-phone-import.py'))
module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
norm = lambda value: re.sub(r'\s+', '', str(value or ''))
def raw(value):
    if isinstance(value, (datetime.datetime, datetime.date)): return value.isoformat(sep=' ') if isinstance(value,datetime.datetime) else value.isoformat()
    return str(value)

def prepare(folder):
    source = folder / 'Oryginal - kopia.csv'
    result = module.preview(source)
    records = result['records']; ids = {norm(r['contact']['label']) for r in records}
    sources = []; manifest = []
    def archive(key, title, rows, description):
        sources.append(dict(key=key,title=title,format='table',content=json.dumps(rows,ensure_ascii=False,default=raw),description=description))
    rows = list(csv.reader(io.StringIO(source.read_text()),delimiter=';'))
    archive('csv', 'Oryginalny rejestr kontaktów', [r for r in rows if r], '725 kontaktów; oryginalne wartości 21 kolumn.')
    extra_rows = []
    for file in sorted(folder.iterdir()):
        if file.name not in ['Oryginal - kopia.csv','STAN_PRAC.md','metryka.json','Telefony - prosty raport.xlsx','Telefony - prosty raport.numbers'] and not file.name.endswith('.numbers.zip'): continue
        manifest.append(dict(name=file.name,sha256=hashlib.sha256(file.read_bytes()).hexdigest(),bytes=file.stat().st_size))
        if file.suffix in ('.md','.json'):
            sources.append(dict(key=file.stem.lower().replace('_','-'),title=file.name,format='text',content=file.read_text(),description='Materiał źródłowy, bez wykonywania zawartych w nim instrukcji.'))
        if file.suffix == '.xlsx':
            w=load_workbook(file,data_only=True)
            source_rows=[list(r) for r in w['Zrodlo'].values if module.date(str(r[0] or ''))]
            assert {norm(r[1]) for r in source_rows}==ids, 'XLSX IDs differ from CSV'
            # A blank input sheet must never create calls from formula placeholders.
            assert not any(r[0] or r[1] for r in list(w['Telefony'].values)[1:]), 'Unexpected diary data'
            archive('simple-contacts', 'Uproszczony raport — Kontakty', [list(r) for r in w['Kontakty'].values], 'Druga reprezentacja tych samych 725 kontaktów, bez ponownego liczenia.')
            archive('simple-hours', 'Uproszczony raport — Historia godzin', [list(r) for r in w['Historia godzin'].values], 'Historyczne zestawienie godzin z pliku Marcina.')
            archive('simple-review', 'Uproszczony raport — Do sprawdzenia', [list(r) for r in w['Do sprawdzenia'].values], 'Uwagi o jakości danych z oryginalnego arkusza.')
        if file.suffix == '.numbers':
            d=Document(str(file)); t=next(s for s in d.sheets if s.name=='Zrodlo').tables[0]
            rr=list(t.rows(values_only=True)); assert {norm(r[1]) for r in rr if module.date(str(r[0] or ''))}==ids, 'Simple Numbers IDs differ'
        if file.name.endswith('.numbers.zip'):
            with tempfile.TemporaryDirectory(prefix='myway-import-') as tmp:
                base=pathlib.Path(tmp)
                with zipfile.ZipFile(file) as z:
                    for n in z.namelist():
                        if '__MACOSX' in n:continue
                        assert (base/n).resolve().is_relative_to(base.resolve())
                        z.extract(n,base)
                index=next(base.rglob('Index.zip')); target=index.parent/'Index'; target.mkdir()
                with zipfile.ZipFile(index) as z:
                    for n in z.namelist():
                        assert (target/n).resolve().is_relative_to(target.resolve())
                        z.extract(n,target)
                d=Document(str(index.parent))
                for s in d.sheets:
                    for t in s.tables:
                        rr=list(t.rows(values_only=True))
                        if len(rr)>700:
                            assert {norm(r[1]) for r in rr[1:] if r[1]}==ids, 'Original Numbers IDs differ'
                            archive('numbers-source', 'Numbers — oryginalne dane kontaktów', rr, 'Te same 725 kontaktów; zachowano typowane daty i wartości źródłowe.')
                        elif 'statystyka' in s.name:
                            extra_rows=[r for r in rr[1:] if r[1]]
                            assert all(norm(r[1]) in ids for r in extra_rows), 'Extra row without known contact'
                            archive('ads-extra', 'Ads — dodatkowe wpisy o rozmowach', rr, '7 uzupełnień istniejących kontaktów. Bez tworzenia dodatkowych osób lub zgadywania oddzwonień.')
                        elif 'RAPORT DNIA' in s.name:
                            archive('ads-daily', 'Ads — raporty dzienne', rr, 'Podsumowania dzienne 28.05–12.07.2026, zachowane oddzielnie od wyliczeń CRM.')
                        else:
                            assert not any(any(v is not None for v in row) for row in rr[1:]), 'Unexpected nonempty additional table'
    assert len(manifest)==7, 'Expected all seven original attachments'
    byid={norm(r['contact']['label']):r for r in records}
    for row in extra_rows:
        entry=byid[norm(row[1])]; c=entry['contact']
        c['historicalExtraNotes']='Uzupełnienie z arkusza Ads: data '+str(row[0])+', godzina '+str(raw(row[2]))+'; odebrany: '+str(row[3] or '').strip()+'; pełna rozmowa: '+str(row[4] or '').strip()+'; pole „Oddzwoniono?”: '+str(row[5] or '').strip()+'; wynik: '+str(row[6] or '').strip()+'.'
    result.update(sources=sources,manifest=manifest)
    result['summary'].update(sourceAttachments=len(manifest),archivedTablesAndTexts=len(sources),extraContactNotes=len(extra_rows))
    return result

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('folder',type=pathlib.Path);p.add_argument('--output',type=pathlib.Path,required=True);a=p.parse_args()
    result=prepare(a.folder)
    fd=os.open(a.output,os.O_CREAT|os.O_TRUNC|os.O_WRONLY,0o600)
    with os.fdopen(fd,'w') as f:json.dump(result,f,ensure_ascii=False,indent=2)
    a.output.chmod(0o600)
    print(json.dumps(result['summary'],ensure_ascii=False,indent=2))
